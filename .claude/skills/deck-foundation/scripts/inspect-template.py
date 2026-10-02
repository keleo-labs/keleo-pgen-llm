#!/usr/bin/env python3
"""Inspect a PowerPoint deck to assess its suitability as a branding template.

Reports slide dimensions, theme colours and fonts, and the slide masters and
layouts (with their placeholders) that a generated deck can populate.

Usage:
    inspect-template.py <deck.pptx> [--json] [--layouts] [--slides]
    inspect-template.py <deck.pptx> --score     # suitability summary only

Exit codes: 0 ok, 1 unreadable/not a pptx, 2 bad arguments.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:  # pragma: no cover - environment guard
    sys.stderr.write("python-pptx is required: pip install python-pptx\n")
    raise SystemExit(1)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

# Theme colour slots in the order PowerPoint presents them.
THEME_COLOR_SLOTS = (
    "dk1", "lt1", "dk2", "lt2",
    "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
)


def emu_to_in(value) -> float | None:
    """Convert an EMU measurement to inches, rounded for display."""
    if value is None:
        return None
    return round(Emu(value).inches, 2)


def _theme_element(prs):
    """Return the <a:theme> root for the first slide master, or None."""
    try:
        part = prs.slide_masters[0].part.part_related_by(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme"
        )
    except (KeyError, IndexError):
        return None
    from lxml import etree

    return etree.fromstring(part.blob)


def extract_theme(prs) -> dict:
    """Pull the colour scheme and font scheme out of the deck's theme part."""
    theme = {"name": None, "colors": {}, "fonts": {}}
    root = _theme_element(prs)
    if root is None:
        return theme

    theme["name"] = root.get("name")

    scheme = root.find(f".//{{{A_NS}}}clrScheme")
    if scheme is not None:
        for slot in THEME_COLOR_SLOTS:
            node = scheme.find(f"{{{A_NS}}}{slot}")
            if node is None:
                continue
            srgb = node.find(f"{{{A_NS}}}srgbClr")
            sysclr = node.find(f"{{{A_NS}}}sysClr")
            if srgb is not None:
                theme["colors"][slot] = "#" + srgb.get("val", "").upper()
            elif sysclr is not None:
                theme["colors"][slot] = "#" + (sysclr.get("lastClr") or "").upper()

    fonts = root.find(f".//{{{A_NS}}}fontScheme")
    if fonts is not None:
        for role, tag in (("heading", "majorFont"), ("body", "minorFont")):
            node = fonts.find(f"{{{A_NS}}}{tag}/{{{A_NS}}}latin")
            if node is not None:
                theme["fonts"][role] = node.get("typeface")

    return theme


def describe_placeholder(shape) -> dict:
    """Summarise one placeholder's identity and geometry."""
    fmt = shape.placeholder_format
    return {
        "idx": fmt.idx,
        "type": str(fmt.type).split(" ")[0] if fmt.type is not None else None,
        "name": shape.name,
        "left_in": emu_to_in(shape.left),
        "top_in": emu_to_in(shape.top),
        "width_in": emu_to_in(shape.width),
        "height_in": emu_to_in(shape.height),
    }


def describe_layout(layout, index: int) -> dict:
    """Summarise one slide layout and the placeholders it offers."""
    return {
        "index": index,
        "name": layout.name,
        "placeholders": [describe_placeholder(ph) for ph in layout.placeholders],
    }


def describe_slides(prs) -> list[dict]:
    """Summarise the actual slides present, to reveal real-world layout usage."""
    slides = []
    for i, slide in enumerate(prs.slides):
        title = None
        if slide.shapes.title is not None:
            title = (slide.shapes.title.text or "").strip()[:80] or None
        slides.append(
            {
                "index": i,
                "layout": slide.slide_layout.name,
                "title": title,
                "shape_count": len(slide.shapes),
                "has_picture": any(s.shape_type == 13 for s in slide.shapes),
            }
        )
    return slides


def inspect(path: Path) -> dict:
    prs = Presentation(str(path))
    masters = []
    for mi, master in enumerate(prs.slide_masters):
        masters.append(
            {
                "index": mi,
                "name": master.name,
                "layouts": [
                    describe_layout(layout, li)
                    for li, layout in enumerate(master.slide_layouts)
                ],
            }
        )

    return {
        "file": str(path),
        "slide_size": {
            "width_in": emu_to_in(prs.slide_width),
            "height_in": emu_to_in(prs.slide_height),
            "aspect": aspect_label(prs.slide_width, prs.slide_height),
        },
        "slide_count": len(prs.slides),
        "theme": extract_theme(prs),
        "masters": masters,
        "slides": describe_slides(prs),
    }


def aspect_label(width, height) -> str:
    """Name the slide aspect ratio, falling back to a numeric ratio."""
    if not width or not height:
        return "unknown"
    ratio = width / height
    for label, value in (("16:9", 16 / 9), ("4:3", 4 / 3), ("16:10", 16 / 10)):
        if abs(ratio - value) < 0.02:
            return label
    return f"{ratio:.2f}:1"


def score(report: dict) -> dict:
    """Judge how usable this deck is as a template donor."""
    layouts = [l for m in report["masters"] for l in m["layouts"]]
    theme = report["theme"]
    named = [l for l in layouts if l["placeholders"]]

    findings = []
    if report["slide_size"]["aspect"] != "16:9":
        findings.append(
            f"aspect is {report['slide_size']['aspect']} — modern decks want 16:9"
        )
    if len(layouts) < 4:
        findings.append(f"only {len(layouts)} layouts — limited structural variety")
    if not theme["colors"]:
        findings.append("no theme colour scheme found")
    if not theme["fonts"]:
        findings.append("no theme font scheme found")
    if not named:
        findings.append("no layouts expose placeholders — cannot be populated reliably")

    return {
        "layout_count": len(layouts),
        "layouts_with_placeholders": len(named),
        "theme_colors": len(theme["colors"]),
        "usable": not findings,
        "findings": findings,
    }


def print_human(report: dict, show_layouts: bool, show_slides: bool) -> None:
    size = report["slide_size"]
    print(f"File:     {report['file']}")
    print(
        f"Size:     {size['width_in']}in x {size['height_in']}in  ({size['aspect']})"
    )
    print(f"Slides:   {report['slide_count']}")

    theme = report["theme"]
    print(f"Theme:    {theme['name'] or '(unnamed)'}")
    if theme["fonts"]:
        pairs = ", ".join(f"{k}={v}" for k, v in theme["fonts"].items())
        print(f"Fonts:    {pairs}")
    if theme["colors"]:
        print("Colors:")
        for slot in THEME_COLOR_SLOTS:
            if slot in theme["colors"]:
                print(f"            {slot:<10} {theme['colors'][slot]}")

    for master in report["masters"]:
        print(f"\nMaster {master['index']}: {master['name']}")
        for layout in master["layouts"]:
            count = len(layout["placeholders"])
            print(f"  [{layout['index']:>2}] {layout['name']}  ({count} placeholders)")
            if show_layouts:
                for ph in layout["placeholders"]:
                    print(
                        f"         idx={ph['idx']:<3} {str(ph['type']):<22}"
                        f" {ph['width_in']}x{ph['height_in']}in"
                        f" @ ({ph['left_in']},{ph['top_in']})"
                    )

    if show_slides and report["slides"]:
        print("\nSlides in deck:")
        for s in report["slides"]:
            title = s["title"] or "(no title)"
            print(f"  [{s['index']:>2}] {s['layout']:<28} {title}")

    verdict = score(report)
    print(
        f"\nTemplate suitability: {'USABLE' if verdict['usable'] else 'NEEDS WORK'}"
        f"  ({verdict['layouts_with_placeholders']}/{verdict['layout_count']}"
        f" layouts have placeholders)"
    )
    for finding in verdict["findings"]:
        print(f"  - {finding}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("deck", type=Path, help="path to a .pptx file")
    parser.add_argument("--json", action="store_true", help="emit JSON")
    parser.add_argument(
        "--layouts", action="store_true", help="show placeholder detail per layout"
    )
    parser.add_argument(
        "--slides", action="store_true", help="list the slides and their layouts"
    )
    parser.add_argument(
        "--score", action="store_true", help="print only the suitability verdict"
    )
    args = parser.parse_args()

    if not args.deck.is_file():
        sys.stderr.write(f"not a file: {args.deck}\n")
        return 1

    try:
        report = inspect(args.deck)
    except Exception as exc:  # noqa: BLE001 - surface any pptx parse failure
        sys.stderr.write(f"could not read {args.deck.name}: {exc}\n")
        return 1

    if args.score:
        verdict = score(report)
        verdict["file"] = str(args.deck)
        print(json.dumps(verdict, indent=2) if args.json else _score_line(verdict))
        return 0

    if args.json:
        print(json.dumps(report, indent=2))
    else:
        print_human(report, args.layouts, args.slides)
    return 0


def _score_line(verdict: dict) -> str:
    status = "USABLE" if verdict["usable"] else "NEEDS WORK"
    base = (
        f"{Path(verdict['file']).name}: {status} — "
        f"{verdict['layouts_with_placeholders']}/{verdict['layout_count']} layouts, "
        f"{verdict['theme_colors']} theme colours"
    )
    if verdict["findings"]:
        base += "\n  " + "\n  ".join(verdict["findings"])
    return base


if __name__ == "__main__":
    raise SystemExit(main())
